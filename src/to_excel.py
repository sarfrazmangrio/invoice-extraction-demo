"""Write the extracted invoices, line items and issues to one Excel workbook."""
from __future__ import annotations

from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from validate import RULES

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
OK_FILL = PatternFill("solid", fgColor="E2EFDA")
REVIEW_FILL = PatternFill("solid", fgColor="FCE4D6")
MONEY = "#,##0.00"


def _date(value):
    try:
        return date.fromisoformat(value) if value else None
    except (TypeError, ValueError):
        return value  # keep the raw text so a bad date stays visible


def _clean(value):
    """Strip characters Excel can't store."""
    return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value


def _append(ws, row):
    ws.append([_clean(v) for v in row])
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str) and cell.value.startswith("="):
            cell.data_type = "s"          # text from an invoice must never run as a formula


def _sheet(wb, title, headers, rows, widths=None, money_cols=(), date_cols=()):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        _append(ws, row)
    ws.freeze_panes = "A2"
    for idx, header in enumerate(headers, start=1):
        letter = get_column_letter(idx)
        lengths = [len(str(header))] + [len(str(r[idx - 1])) for r in rows if r[idx - 1] is not None]
        ws.column_dimensions[letter].width = (widths or {}).get(header) or min(max(lengths) + 2, 60)
        if header in money_cols:
            for cell in ws[letter][1:]:
                cell.number_format = MONEY
        if header in date_cols:
            for cell in ws[letter][1:]:
                cell.number_format = "yyyy-mm-dd"
    ws.auto_filter.ref = ws.dimensions
    return ws


def write_workbook(path: Path, records: list[dict], issues: list[dict], failures: list[dict], stats: dict) -> Path:
    """Write the workbook and return where it was saved."""
    by_file = Counter(i["file"] for i in issues)
    wb = Workbook()
    wb.remove(wb.active)

    # Summary
    ws = wb.create_sheet("Summary")
    needs_review = sum(1 for r in records if by_file[r["file"]])
    summary = [
        ("Invoices processed", len(records) + len(failures)),
        ("Extracted", len(records)),
        ("Failed to extract", len(failures)),
        ("OK (no issues)", len(records) - needs_review),
        ("Needs review", needs_review),
        ("Issues found", len(issues)),
        ("Model", stats.get("model")),
        ("Input tokens", stats.get("input_tokens")),
        ("Output tokens", stats.get("output_tokens")),
        ("Cost (USD)", stats.get("cost_usd")),
    ]
    ws.append(["Item", "Value"])
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for row in summary:
        _append(ws, list(row))
    ws.append([])
    ws.append(["Issues by rule", "Count"])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    rule_counts = Counter(i["rule"] for i in issues)
    for rule, label in RULES.items():
        if rule_counts[rule]:
            ws.append([label, rule_counts[rule]])
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 28

    # Invoices
    rows = []
    for r in records:
        d = r["data"]
        rows.append([
            r["file"], "Needs review" if by_file[r["file"]] else "OK", by_file[r["file"]],
            d.get("vendor_name"), d.get("invoice_number"), _date(d.get("invoice_date")), _date(d.get("due_date")),
            d.get("customer_name"), d.get("currency"), len(d.get("line_items") or []),
            d.get("subtotal"), d.get("tax_rate_percent"), d.get("tax_amount"), d.get("total"),
        ])
    inv = _sheet(wb, "Invoices",
                 ["File", "Status", "Issues", "Vendor", "Invoice no.", "Invoice date", "Due date", "Customer",
                  "Currency", "Lines", "Subtotal", "Tax rate %", "Tax", "Total"],
                 rows, money_cols=("Subtotal", "Tax", "Total"), date_cols=("Invoice date", "Due date"))
    for row in inv.iter_rows(min_row=2):
        row[1].fill = REVIEW_FILL if row[1].value == "Needs review" else OK_FILL

    # Line items
    rows = []
    for r in records:
        d = r["data"]
        for n, line in enumerate(d.get("line_items") or [], start=1):
            rows.append([r["file"], d.get("invoice_number"), n, line.get("description"),
                         line.get("quantity"), line.get("unit_price"), line.get("amount")])
    _sheet(wb, "Line items", ["File", "Invoice no.", "Line", "Description", "Quantity", "Unit price", "Amount"],
           rows, money_cols=("Unit price", "Amount"))

    # Issues
    rows = [[i["file"], i["vendor_name"], i["invoice_number"], RULES.get(i["rule"], i["rule"]), i["severity"],
             i["message"], i["expected"], i["found"]] for i in issues]
    rows += [[f["file"], None, None, "Extraction failed", "error", f["error"], None, None] for f in failures]
    _sheet(wb, "Issues", ["File", "Vendor", "Invoice no.", "Check", "Severity", "Detail", "Expected", "Found"], rows)

    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(path)
    except PermissionError:  # usually the file is open in Excel
        path = path.with_name(f"{path.stem}_{datetime.now():%Y%m%d_%H%M%S}{path.suffix}")
        wb.save(path)
        print(f"The workbook was open in another program, so this run saved to {path.name}")
    return path
